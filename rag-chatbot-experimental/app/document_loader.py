from __future__ import annotations

import logging
import hashlib
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

from app.schemas import DocumentPart


logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class DocumentLoadingError(RuntimeError):
    """Raised when a document cannot be read."""


class DocumentLoader:
    def __init__(self, documents_dir: Path) -> None:
        self.documents_dir = documents_dir

    def iter_files(self) -> Iterable[Path]:
        if not self.documents_dir.exists():
            return []
        return sorted(
            path
            for path in self.documents_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
        )

    def load_all(self) -> tuple[list[DocumentPart], list[str]]:
        files = list(self.iter_files())
        if not files:
            logger.warning("Aucun document supporté trouvé dans %s", self.documents_dir)
            return [], []

        parts: list[DocumentPart] = []
        failed_files: list[str] = []
        for path in files:
            try:
                loaded_parts = self.load_file(path)
                parts.extend(loaded_parts)
                logger.info("Document chargé: %s (%s partie(s))", path.name, len(loaded_parts))
            except Exception as exc:
                logger.warning("Document ignoré car illisible: %s (%s)", path.name, exc)
                failed_files.append(str(path))
        return parts, failed_files

    def load_file(self, path: Path) -> list[DocumentPart]:
        suffix = path.suffix.lower()
        ingested_at = datetime.now(UTC).isoformat()
        resolved_path = path.resolve()
        base_metadata = {
            "document_id": hashlib.sha256(str(resolved_path).encode("utf-8")).hexdigest()[:24],
            "file_name": path.name,
            "file_path": str(resolved_path),
            "source": str(resolved_path),
            "source_path": str(resolved_path),
            "source_type": suffix.lstrip("."),
            "file_type": suffix.lstrip("."),
            "file_hash": _file_hash(path),
            "last_modified": datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).isoformat(),
            "indexed_at": ingested_at,
            "ingested_at": ingested_at,
            "allowed_groups": ["public"],
        }

        if suffix in {".txt", ".md"}:
            return self._load_text_file(path, base_metadata)
        if suffix == ".pdf":
            return self._load_pdf(path, base_metadata)
        if suffix == ".docx":
            return self._load_docx(path, base_metadata)

        raise DocumentLoadingError(f"Extension non supportée: {suffix}")

    def _load_text_file(self, path: Path, metadata: dict[str, object]) -> list[DocumentPart]:
        text = path.read_text(encoding="utf-8", errors="replace")
        text = _clean_text(text)
        section_title = _extract_section_title(text, markdown=path.suffix.lower() == ".md")
        return [DocumentPart(text=text, metadata={**metadata, "page": None, "section_title": section_title})] if text else []

    def _load_pdf(self, path: Path, metadata: dict[str, object]) -> list[DocumentPart]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentLoadingError("Le paquet `pypdf` est requis pour lire les PDF.") from exc

        reader = PdfReader(str(path))
        parts: list[DocumentPart] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = _clean_text(page.extract_text() or "")
            if text:
                parts.append(DocumentPart(text=text, metadata={**metadata, "page": page_index, "section_title": None}))
        if not parts:
            raise DocumentLoadingError(
                "Aucun texte extractible trouvé. Le PDF est probablement scanné ou nécessite un OCR."
            )
        return parts

    def _load_docx(self, path: Path, metadata: dict[str, object]) -> list[DocumentPart]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentLoadingError("Le paquet `python-docx` est requis pour lire les DOCX.") from exc

        document = DocxDocument(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = _clean_text("\n\n".join(paragraphs))
        return [
            DocumentPart(text=text, metadata={**metadata, "page": None, "section_title": _extract_section_title(text)})
        ] if text else []


def _clean_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized.split("\n")]
    cleaned_lines: list[str] = []
    blank_seen = False
    for line in lines:
        if not line:
            if not blank_seen:
                cleaned_lines.append("")
            blank_seen = True
            continue
        cleaned_lines.append(line)
        blank_seen = False
    return "\n".join(cleaned_lines).strip()


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as document:
        for block in iter(lambda: document.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _extract_section_title(text: str, *, markdown: bool = False) -> str | None:
    for line in text.splitlines():
        cleaned = line.strip()
        if not cleaned:
            continue
        if markdown:
            heading = re.match(r"^#{1,6}\s+(.+)$", cleaned)
            if heading:
                return heading.group(1).strip()[:200]
        if len(cleaned) <= 120:
            return cleaned.lstrip("#").strip()[:200]
        break
    return None
